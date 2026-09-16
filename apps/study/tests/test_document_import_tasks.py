import pytest
from unittest.mock import Mock
from docx import Document
from PIL import Image
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from apps.accounts.models import UserAccount
from apps.common.ai.exceptions import AIResponseError
from apps.courses.models import Course
from apps.study import document_import_tasks as module
from apps.study.document_import_models import QuestionDocumentImportItem, QuestionDocumentImportTask
from apps.study.document_import_service import ExtractedDocument, ExtractedQuestionFragment
from apps.study.models import QuestionIngestionBatch


@pytest.mark.django_db
def test_document_task_records_each_source_position_when_two_fragments_share_one_question(task, monkeypatch):
    """Repeated document labels retain two source references after content deduplication."""
    from apps.courses.models import CourseQuestionDocumentReference

    monkeypatch.setattr(
        module,
        'extract_document',
        lambda _path: ExtractedDocument(
            document_type='pdf', page_count=1,
            fragments=[
                ExtractedQuestionFragment(question_no='1', text='1. First copy.', page_range=(1, 1)),
                ExtractedQuestionFragment(question_no='1', text='1. Second copy.', page_range=(1, 1)),
            ],
        ),
    )
    monkeypatch.setattr(module, 'structure_candidate', Mock(return_value=valid_question()))

    module.process_document_import_task.run(str(task.id))

    references = CourseQuestionDocumentReference.objects.filter(document_import_task=task).order_by('source_position')
    assert references.count() == 2
    assert list(references.values_list('source_position', 'source_question_no')) == [(1, '1'), (2, '1')]


@pytest.mark.django_db
def test_async_document_task_persists_and_finalizes_question_items(task, monkeypatch, settings, tmp_path):
    """The production task path completes through durable item tasks, not the legacy fallback."""
    settings.CELERY_TASK_ALWAYS_EAGER = True
    source_path = tmp_path / 'questions.pdf'
    source_path.write_bytes(b'async document fixture')
    task.source_file = str(source_path)
    task.save(update_fields=['source_file'])
    monkeypatch.setattr(
        module,
        'extract_document',
        lambda _path: ExtractedDocument(
            document_type='pdf', page_count=1,
            fragments=[
                ExtractedQuestionFragment(question_no='1', text='1. First.'),
                ExtractedQuestionFragment(question_no='2', text='2. Second.'),
            ],
        ),
    )
    first, second = valid_question(), valid_question()
    second['question_no'] = '2'
    monkeypatch.setattr(module, 'structure_candidate', Mock(side_effect=[first, second]))
    result = module.process_document_import_task.apply_async(args=(str(task.id),))

    assert result.successful()
    task.refresh_from_db()
    assert task.stage == QuestionDocumentImportTask.Stage.SUCCESS
    assert task.candidate_count == 2
    assert QuestionDocumentImportItem.objects.filter(
        document_import_task=task,
        status=QuestionDocumentImportItem.Status.SUCCESS,
        ingest_status=QuestionDocumentImportItem.IngestStatus.INGESTED,
    ).count() == 2


@pytest.mark.django_db
def test_document_task_keeps_duplicate_labels_bound_to_distinct_content_fingerprints(task, monkeypatch):
    """A reset label must never collapse two different original questions."""
    from apps.courses.models import CourseQuestionDocumentReference

    monkeypatch.setattr(
        module,
        'extract_document',
        lambda _path: ExtractedDocument(
            document_type='docx', page_count=2,
            fragments=[
                ExtractedQuestionFragment(
                    question_no='1', text='1. First original question.',
                    page_range=(1, 1), source_paragraph_index=3, section_path='section-1',
                ),
                ExtractedQuestionFragment(
                    question_no='1', text='1. Second original question.',
                    page_range=(2, 2), source_paragraph_index=21, section_path='section-2',
                ),
            ],
        ),
    )
    first, second = valid_question(), valid_question()
    first['stem'], second['stem'] = 'First structured stem.', 'Second structured stem.'
    monkeypatch.setattr(module, 'structure_candidate', Mock(side_effect=[first, second]))

    module.process_document_import_task.run(str(task.id))

    references = list(CourseQuestionDocumentReference.objects.filter(
        document_import_task=task,
    ).order_by('source_position'))
    assert len(references) == 2
    assert len({reference.course_question_link.question_id for reference in references}) == 2
    assert [reference.source_locator for reference in references] == [
        'section-1/paragraph-3/pages-1-1/question-1',
        'section-2/paragraph-21/pages-2-2/question-1',
    ]


@pytest.mark.django_db
def test_reprocessing_rebinds_a_stale_source_fingerprint_by_recomputed_content(task, monkeypatch):
    """Historical raw-source bindings are corrected instead of trusted blindly."""
    from apps.parser.models import QuestionDocumentSourceFingerprint

    fragment = ExtractedQuestionFragment(question_no='1', text='1. Same original text.')
    monkeypatch.setattr(
        module, 'extract_document',
        lambda _path: ExtractedDocument(document_type='pdf', page_count=1, fragments=[fragment]),
    )
    stale = valid_question()
    stale['stem'] = 'Stale structured stem.'
    monkeypatch.setattr(module, 'structure_candidate', Mock(return_value=stale))
    module.process_document_import_task.run(str(task.id))
    source_fingerprint = module.document_source_fingerprint(fragment)
    stale_question_id = QuestionDocumentSourceFingerprint.objects.get(
        fingerprint=source_fingerprint,
    ).canonical_question_id

    retry_batch = QuestionIngestionBatch.objects.create(
        actor=task.batch.actor,
        course=task.course,
        source_type=QuestionIngestionBatch.SourceType.DOCUMENT_IMPORT,
        source_name='questions-retry.pdf',
    )
    retry_task = QuestionDocumentImportTask.objects.create(
        batch=retry_batch, course=task.course, source_file='imports/questions-retry.pdf',
        detected_mime='application/pdf', document_type='pdf', page_count=1,
    )
    corrected = valid_question()
    corrected['stem'] = 'Correct structured stem.'
    monkeypatch.setattr(module, 'structure_candidate', Mock(return_value=corrected))
    module.process_document_import_task.run(str(retry_task.id))

    repaired = QuestionDocumentSourceFingerprint.objects.get(fingerprint=source_fingerprint)
    assert repaired.canonical_question_id != stale_question_id


def valid_question():
    return {
        'question_no': '1',
        'question_type': 'single_choice',
        'stem': 'Choose one.',
        'options': [{'label': 'A', 'content': 'One'}],
        'answer': 'A',
        'analysis': 'Only A is one.',
        'tables': [],
        'illustrations': [],
        'confidence': 0.9,
        'review_reason': '',
        'source': {'page_start': 1, 'page_end': 1},
    }


def _asset_question(question_no, asset_file):
    question = valid_question()
    question['question_no'] = question_no
    question['stem'] = f'Question {question_no} has a diagram.'
    question['illustrations'] = [{'file': asset_file}]
    return question


def _pdf_with_image(tmp_path):
    image_path = tmp_path / 'pdf-diagram.png'
    Image.new('RGB', (10, 10), color='black').save(image_path)
    document_path = tmp_path / 'questions.pdf'
    pdf = canvas.Canvas(str(document_path), pagesize=letter)
    pdf.drawString(72, 720, '1. PDF diagram question')
    pdf.drawImage(str(image_path), 72, 650, width=20, height=20)
    pdf.save()
    return document_path


def _docx_with_image(tmp_path):
    image_path = tmp_path / 'docx-diagram.png'
    Image.new('RGB', (10, 10), color='black').save(image_path)
    document_path = tmp_path / 'questions.docx'
    document = Document()
    document.add_paragraph('1. DOCX diagram question')
    document.add_picture(str(image_path))
    document.save(document_path)
    return document_path


@pytest.fixture
def task(db):
    teacher = UserAccount.objects.create(
        mobile='13900008302', display_name='Document task teacher', role_type='teacher',
    )
    course = Course.objects.create(
        name='Document task course', subject='math', grade_level='Grade 8', teacher=teacher,
    )
    batch = QuestionIngestionBatch.objects.create(
        actor=teacher,
        course=course,
        source_type=QuestionIngestionBatch.SourceType.DOCUMENT_IMPORT,
        source_name='questions.pdf',
    )
    return QuestionDocumentImportTask.objects.create(
        batch=batch,
        course=course,
        source_file='imports/questions.pdf',
        detected_mime='application/pdf',
        document_type='pdf',
        page_count=2,
    )


@pytest.mark.django_db
def test_task_marks_partial_success_after_one_candidate_failure(task, monkeypatch):
    """One malformed Qwen candidate never prevents a valid sibling from importing."""
    monkeypatch.setattr(
        module,
        'extract_document',
        lambda _path: ExtractedDocument(
            document_type='pdf', page_count=2,
            fragments=[
                ExtractedQuestionFragment(question_no='1', text='1. Choose one.'),
                ExtractedQuestionFragment(question_no='2', text='2. Broken candidate.'),
            ],
        ),
    )
    monkeypatch.setattr(
        module,
        'structure_candidate',
        Mock(side_effect=[valid_question(), AIResponseError('bad json')]),
    )

    module.process_document_import_task.run(str(task.id))

    task.refresh_from_db()
    assert task.stage == 'partial_success'


@pytest.mark.parametrize(
    ('question_no', 'message', 'expected'),
    [
        ('12', 'AI provider request timed out', '第12题：AI响应超时'),
        ('18', 'AI response is not valid JSON; preview={"question_no":"18"}', '第18题：AI返回格式错误'),
    ],
)
def test_question_error_is_safe_and_readable_for_history_popup(question_no, message, expected):
    assert module._question_error(question_no, ValueError(message)) == expected


@pytest.mark.django_db
@pytest.mark.parametrize('factory', [_pdf_with_image, _docx_with_image])
def test_document_asset_is_materialized_before_shared_ingestion(task, tmp_path, monkeypatch, factory):
    """PDF and DOCX internal images become readable source-root files and QuestionImages."""
    source_path = factory(tmp_path)
    task.source_file = str(source_path)
    task.document_type = source_path.suffix.lstrip('.')
    task.save(update_fields=['source_file', 'document_type'])

    def structure(candidate, *, asset_files=None):
        assert asset_files
        question = _asset_question(candidate.question_no, next(iter(asset_files.values())))
        source_asset = candidate.asset_refs[0]
        question['illustrations'][0].update({
            'placement': 'options',
            'bbox': list(source_asset.bbox) if source_asset.bbox else None,
            'source_page': source_asset.page_no,
        })
        return question

    monkeypatch.setattr(module, 'structure_candidate', structure)

    module.process_document_import_task.run(str(task.id))

    task.refresh_from_db()
    from apps.parser.models import QuestionImage
    image = QuestionImage.objects.get()
    assert task.stage == 'success'
    assert image.file_path
    assert image.original_file_path
    assert image.placement == 'options'
    if source_path.suffix.lower() == '.pdf':
        assert image.bbox


@pytest.mark.django_db
def test_task_continues_after_exhausted_asset_read_retry(task, monkeypatch):
    """A candidate that still raises OSError after its own retry does not stop later candidates."""
    monkeypatch.setattr(
        module,
        'extract_document',
        lambda _path: ExtractedDocument(
            document_type='pdf', page_count=3,
            fragments=[
                ExtractedQuestionFragment(question_no='1', text='1. First.'),
                ExtractedQuestionFragment(question_no='2', text='2. Asset failure.'),
                ExtractedQuestionFragment(question_no='3', text='3. Last.'),
            ],
        ),
    )
    third = valid_question()
    third['question_no'] = '3'
    monkeypatch.setattr(
        module,
        'structure_candidate',
        Mock(side_effect=[valid_question(), OSError(r'C:\private\asset.png unavailable'), third]),
    )

    module.process_document_import_task.run(str(task.id))

    task.refresh_from_db()
    assert task.stage == 'partial_success'
    assert task.progress == 100
    assert '第2题：题目资源处理失败' in task.error_summary
