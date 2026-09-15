from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from apps.study.document_import_service import (
    DocumentValidationError,
    extract_document,
    validate_document_upload,
)


def make_upload(name, content, content_type):
    return SimpleUploadedFile(name, content, content_type=content_type)


def _over_limit_text():
    return ''.join(f'{value:06x}' for value in range(26_667))


class SizeLyingUpload:
    """A stream whose metadata says 1 byte but that contains over 100 MiB."""

    name = 'stream.pdf'
    size = 1

    def __init__(self):
        self._position = 0
        self._length = 100 * 1024 * 1024 + 1

    def seek(self, position):
        self._position = position

    def read(self, size=-1):
        if size < 0:
            raise AssertionError('upload validation must use bounded reads')
        if self._position >= self._length:
            return b''
        count = min(size, self._length - self._position)
        start = self._position
        self._position += count
        if start == 0:
            return b'%PDF-1.7' + b'x' * (count - len(b'%PDF-1.7'))
        return b'x' * count


def _repeated_image_docx(tmp_path):
    image_path = tmp_path / 'reused.png'
    Image.new('RGB', (10, 10), color='black').save(image_path)
    document_path = tmp_path / 'repeated-image.docx'
    document = Document()
    document.add_paragraph('1. Repeated image question')
    for _ in range(101):
        document.add_picture(str(image_path))
    document.save(document_path)
    with ZipFile(document_path) as archive:
        assert archive.read('word/document.xml').count(b'r:embed=') == 101
    return document_path


def _compressed_docx_bomb(tmp_path):
    source_path = tmp_path / 'source.docx'
    document = Document()
    document.add_paragraph('1. Safe question shell')
    document.save(source_path)
    bomb_path = tmp_path / 'compressed-bomb.docx'
    with ZipFile(source_path) as source, ZipFile(bomb_path, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == 'word/document.xml':
                payload = payload.replace(
                    b'</w:body>', b'<!--' + b'A' * (3 * 1024 * 1024) + b'--></w:body>',
                )
            target.writestr(info.filename, payload)
    assert bomb_path.stat().st_size < 100 * 1024 * 1024
    return bomb_path


def _nested_table_over_limit_docx(tmp_path):
    document_path = tmp_path / 'nested-table-over-limit.docx'
    document = Document()
    document.add_paragraph('1. Nested table question')
    outer_table = document.add_table(rows=1, cols=1)
    nested_table = outer_table.cell(0, 0).add_table(rows=1, cols=1)
    nested_table.cell(0, 0).text = _over_limit_text()
    document.save(document_path)
    return document_path


def _set_word_auto_number(paragraph, number_id='1', level='0'):
    """Add the numbering metadata Word uses for a visible automatic list marker."""
    number_properties = OxmlElement('w:numPr')
    indentation_level = OxmlElement('w:ilvl')
    indentation_level.set(qn('w:val'), level)
    number = OxmlElement('w:numId')
    number.set(qn('w:val'), number_id)
    number_properties.extend([indentation_level, number])
    paragraph._p.get_or_add_pPr().append(number_properties)


def _docx_with_safe_member_count(tmp_path, member_count):
    source_path = tmp_path / 'source.docx'
    document = Document()
    document.add_paragraph('1. Safe question shell')
    document.save(source_path)

    document_path = tmp_path / 'many-members.docx'
    with ZipFile(source_path) as source, ZipFile(document_path, 'w', ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, source.read(info.filename))
        for index in range(member_count - len(source.infolist())):
            target.writestr(f'word/media/safe-{index}.bin', b'safe')
    return document_path


def test_validate_document_rejects_docx_named_pdf():
    """Changing a PDF filename to .docx must not bypass type validation."""
    upload = make_upload('fake.docx', b'%PDF-1.7', 'application/pdf')

    with pytest.raises(DocumentValidationError, match='文件内容与扩展名不一致'):
        validate_document_upload(upload)


def test_validate_document_rejects_legacy_doc_extension():
    """Legacy Word binaries are outside the deliberately small import surface."""
    upload = make_upload('legacy.doc', b'\xd0\xcf\x11\xe0', 'application/msword')

    with pytest.raises(DocumentValidationError, match='仅支持 PDF 和 DOCX'):
        validate_document_upload(upload)


def test_validate_document_rejects_file_larger_than_100_mb():
    """The declared upload size is checked before expensive document parsing."""
    upload = make_upload('large.pdf', b'%PDF-1.7', 'application/pdf')
    upload.size = 100 * 1024 * 1024 + 1

    with pytest.raises(DocumentValidationError, match='100MB'):
        validate_document_upload(upload)


def test_validate_document_rejects_pdf_over_page_limit(tmp_path):
    """An over-limit PDF is rejected at upload validation, before it is queued."""
    pdf_path = tmp_path / 'too-many-pages.pdf'
    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    for _ in range(101):
        pdf.drawString(72, 720, '1. A numbered question')
        pdf.showPage()
    pdf.save()

    with pytest.raises(DocumentValidationError, match='100 页'):
        validate_document_upload(make_upload('too-many-pages.pdf', pdf_path.read_bytes(), 'application/pdf'))


def test_validate_document_rejects_docx_over_equivalent_page_limit(tmp_path):
    """The DOCX complexity formula is also enforced before queueing work."""
    document_path = tmp_path / 'too-large.docx'
    document = Document()
    document.add_paragraph('1. ' + _over_limit_text())
    document.save(document_path)

    with pytest.raises(DocumentValidationError, match='100 页'):
        validate_document_upload(make_upload(
            'too-large.docx', document_path.read_bytes(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        ))


def test_validate_document_rejects_size_lying_stream_over_100_mb():
    """The upload byte stream, not caller-controlled metadata, enforces the size cap."""
    with pytest.raises(DocumentValidationError, match='100MB'):
        validate_document_upload(SizeLyingUpload())


def test_validate_document_counts_each_reused_docx_image_reference(tmp_path):
    """Every displayed image counts even when the DOCX reuses one relationship target."""
    document_path = _repeated_image_docx(tmp_path)

    with pytest.raises(DocumentValidationError, match='100 页'):
        validate_document_upload(make_upload(
            'repeated-image.docx', document_path.read_bytes(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        ))


def test_extract_document_counts_each_reused_docx_image_reference(tmp_path):
    """Path-based extraction applies the same repeated-image quota as upload validation."""
    with pytest.raises(DocumentValidationError, match='100 页'):
        extract_document(_repeated_image_docx(tmp_path))


def test_validate_document_rejects_docx_compression_bomb(tmp_path):
    """A valid DOCX ZIP with a dangerous expansion ratio is rejected before parsing."""
    bomb_path = _compressed_docx_bomb(tmp_path)

    with pytest.raises(DocumentValidationError, match='压缩比'):
        validate_document_upload(make_upload(
            'compressed-bomb.docx', bomb_path.read_bytes(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        ))


def test_extract_document_rejects_docx_compression_bomb(tmp_path):
    """Path-based DOCX extraction rejects a dangerous ZIP before python-docx opens it."""
    with pytest.raises(DocumentValidationError, match='压缩比'):
        extract_document(_compressed_docx_bomb(tmp_path))


def test_validate_document_counts_nested_table_text_and_table_count(tmp_path):
    """Nested table content contributes to the upload-time DOCX equivalent page limit."""
    document_path = _nested_table_over_limit_docx(tmp_path)

    with pytest.raises(DocumentValidationError, match='100 页'):
        validate_document_upload(make_upload(
            'nested-table-over-limit.docx', document_path.read_bytes(),
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        ))


def test_extract_document_counts_nested_table_text_and_table_count(tmp_path):
    """Nested table content contributes to the path-based DOCX equivalent page limit."""
    with pytest.raises(DocumentValidationError, match='100 页'):
        extract_document(_nested_table_over_limit_docx(tmp_path))


@pytest.fixture
def pdf_fixture(tmp_path):
    image_path = tmp_path / 'diagram.png'
    image = Image.new('RGB', (10, 10), color='black')
    image.save(image_path)

    pdf_path = tmp_path / 'questions.pdf'
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.drawString(72, 720, '1. First question has an illustration.')
    pdf.drawImage(str(image_path), 72, 650, width=20, height=20)
    pdf.drawString(72, 580, '2. Second question has no illustration.')
    pdf.save()
    pdf_path.write_bytes(buffer.getvalue())
    return pdf_path


def test_extract_pdf_slices_questions_and_binds_assets(pdf_fixture):
    """A numbered PDF is split into independently processable question fragments."""
    document = extract_document(pdf_fixture)

    assert [item.question_no for item in document.fragments] == ['1', '2']
    assert document.fragments[0].asset_refs
    assert document.fragments[0].page_range == (1, 1)


def test_extract_pdf_rejects_more_than_100_pages(tmp_path):
    """Page limits are enforced before a large PDF can enter AI processing."""
    pdf_path = tmp_path / 'too-many-pages.pdf'
    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    for _ in range(101):
        pdf.drawString(72, 720, '1. A numbered question')
        pdf.showPage()
    pdf.save()

    with pytest.raises(DocumentValidationError, match='100 页'):
        extract_document(pdf_path)


def test_extract_document_rejects_without_numbered_candidates(tmp_path):
    """Unstructured documents stop before any downstream AI request is possible."""
    pdf_path = tmp_path / 'no-questions.pdf'
    pdf = canvas.Canvas(str(pdf_path), pagesize=letter)
    pdf.drawString(72, 720, 'This page has no numbered question candidates.')
    pdf.save()

    with pytest.raises(DocumentValidationError, match='未找到可识别的编号题目'):
        extract_document(pdf_path)


def test_extract_docx_rejects_equivalent_page_count_above_limit(tmp_path):
    """DOCX complexity is limited by the specified text/table/image page estimate."""
    document_path = tmp_path / 'too-large.docx'
    document = Document()
    document.add_paragraph('1. ' + _over_limit_text())
    document.save(document_path)

    with pytest.raises(DocumentValidationError, match='100 页'):
        extract_document(document_path)


def test_extract_docx_slices_questions_and_attaches_tables_and_images(tmp_path):
    """DOCX paragraphs, nearby tables, and relationship images stay with their question."""
    image_path = tmp_path / 'diagram.png'
    Image.new('RGB', (10, 10), color='black').save(image_path)
    document_path = tmp_path / 'questions.docx'
    document = Document()
    document.add_paragraph('1. First DOCX question')
    document.add_picture(str(image_path))
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Given value'
    document.add_paragraph('2. Second DOCX question')
    document.save(document_path)

    extracted = extract_document(document_path)

    assert [fragment.question_no for fragment in extracted.fragments] == ['1', '2']
    assert extracted.fragments[0].tables == [[['Given value']]]
    assert extracted.fragments[0].asset_refs


def test_validate_document_accepts_safe_docx_with_3000_members(tmp_path):
    """A normal image-heavy Word paper must reach parsing while byte safety limits still apply."""
    document_path = _docx_with_safe_member_count(tmp_path, 3000)

    validated = validate_document_upload(make_upload(
        'many-members.docx', document_path.read_bytes(),
        'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    ))

    assert validated.document_type == 'docx'


def test_extract_docx_slices_word_automatic_numbered_questions(tmp_path):
    """Visible Word list numbers, although absent from paragraph.text, delimit question fragments."""
    document_path = tmp_path / 'automatic-numbering.docx'
    document = Document()
    first = document.add_paragraph('First automatically numbered question')
    _set_word_auto_number(first)
    document.add_paragraph('A. First option')
    document.add_paragraph('B. Second option')
    second = document.add_paragraph('Second automatically numbered question')
    _set_word_auto_number(second)
    document.add_paragraph('A. Third option')
    document.save(document_path)

    extracted = extract_document(document_path)

    assert [fragment.question_no for fragment in extracted.fragments] == ['1', '2']
    assert extracted.fragments[0].text.startswith('1. First automatically numbered question')
    assert 'B. Second option' in extracted.fragments[0].text


def test_extract_docx_falls_back_to_blank_separated_option_groups(tmp_path):
    """Unnumbered practice sheets still yield separate candidate questions for AI structuring."""
    document_path = tmp_path / 'unnumbered-practice.docx'
    document = Document()
    document.add_paragraph('Lesson heading')
    document.add_paragraph('')
    document.add_paragraph('First unnumbered question')
    document.add_paragraph('A. First option')
    document.add_paragraph('B. Second option')
    document.add_paragraph('')
    document.add_paragraph('Second unnumbered question')
    document.add_paragraph('A. Third option')
    document.add_paragraph('B. Fourth option')
    document.save(document_path)

    extracted = extract_document(document_path)

    assert [fragment.question_no for fragment in extracted.fragments] == ['1', '2']
    assert extracted.fragments[0].text == 'First unnumbered question\nA. First option\nB. Second option'
    assert extracted.fragments[1].text == 'Second unnumbered question\nA. Third option\nB. Fourth option'
