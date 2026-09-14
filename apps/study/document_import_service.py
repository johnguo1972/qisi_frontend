"""Safe, deterministic extraction primitives for course PDF and DOCX imports."""

from dataclasses import dataclass, field
from io import BytesIO
from math import ceil
from pathlib import Path
import re
from zipfile import BadZipFile, LargeZipFile, ZipFile

import fitz
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


MAX_DOCUMENT_BYTES = 100 * 1024 * 1024
MAX_DOCUMENT_PAGES = 100
READ_CHUNK_BYTES = 1024 * 1024
MAX_DOCX_ARCHIVE_MEMBERS = 3000
MAX_DOCX_MEMBER_BYTES = 25 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCX_COMPRESSION_RATIO = 200
QUESTION_START = re.compile(r'(?m)^\s*(\d+)\s*[.、．)]\s*')
OPTION_START = re.compile(r'(?m)^\s*[A-HＡ-Ｈ]\s*[.．、]\s*')
QUESTION_SIGNAL = re.compile(r'[？?（(＿_]|填空|简答|计算|证明|实验|作图')
RELATION_IMAGE_SUFFIX = '/image'
DOCX_IMAGE_REFERENCE = re.compile(r'r:embed="([^"]+)"')


class DocumentValidationError(ValueError):
    """Raised before an invalid document can reach an AI processing stage."""


@dataclass(frozen=True)
class ValidatedDocument:
    filename: str
    document_type: str
    detected_mime: str
    size_bytes: int
    page_count: int


@dataclass(frozen=True)
class DocumentAssetRef:
    reference: str
    page_no: int | None = None
    bbox: tuple[float, float, float, float] | None = None


@dataclass
class ExtractedQuestionFragment:
    question_no: str
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)
    page_range: tuple[int, int] = (1, 1)
    asset_refs: list[DocumentAssetRef] = field(default_factory=list)


@dataclass
class ExtractedDocument:
    document_type: str
    page_count: int
    fragments: list[ExtractedQuestionFragment]


def validate_document_upload(uploaded_file) -> ValidatedDocument:
    """Validate a Django upload using both its extension and its actual bytes."""
    filename = str(getattr(uploaded_file, 'name', ''))
    extension = Path(filename).suffix.lower()
    if extension not in {'.pdf', '.docx'}:
        raise DocumentValidationError('仅支持 PDF 和 DOCX 格式文件')

    size_bytes = int(getattr(uploaded_file, 'size', 0))
    if size_bytes > MAX_DOCUMENT_BYTES:
        raise DocumentValidationError('文件大小不能超过 100MB')

    content = _read_upload_bytes(uploaded_file)
    document_type = _detect_document_type(content)
    if document_type is None:
        raise DocumentValidationError('文件内容无法识别为 PDF 或 DOCX')
    if document_type != extension.lstrip('.'):
        raise DocumentValidationError('文件内容与扩展名不一致')
    page_count = _validate_document_page_limit(document_type, content)
    return ValidatedDocument(
        filename=filename,
        document_type=document_type,
        detected_mime='application/pdf' if document_type == 'pdf'
        else 'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        size_bytes=size_bytes,
        page_count=page_count,
    )


def extract_document(source_path) -> ExtractedDocument:
    """Extract numbered question candidates and nearby image references from a document."""
    path = Path(source_path)
    _validate_source_path(path)
    document_type = path.suffix.lower().lstrip('.')
    if document_type == 'pdf':
        return _extract_pdf(path)
    return _extract_docx(path)


def _read_upload_bytes(uploaded_file) -> bytes:
    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass
    content = _read_limited_bytes(uploaded_file)
    try:
        uploaded_file.seek(0)
    except (AttributeError, OSError):
        pass
    return content


def _read_limited_bytes(stream) -> bytes:
    """Read no more than the real upload limit, independent of metadata."""
    chunks = []
    bytes_read = 0
    maximum_to_read = MAX_DOCUMENT_BYTES + 1
    while bytes_read < maximum_to_read:
        chunk = stream.read(min(READ_CHUNK_BYTES, maximum_to_read - bytes_read))
        if not chunk:
            break
        bytes_read += len(chunk)
        if bytes_read > MAX_DOCUMENT_BYTES:
            raise DocumentValidationError('文件大小不能超过 100MB')
        chunks.append(chunk)
    return b''.join(chunks)


def _detect_document_type(content: bytes) -> str | None:
    if content.startswith(b'%PDF-'):
        return 'pdf'
    try:
        with ZipFile(BytesIO(content)) as archive:
            if 'word/document.xml' in archive.namelist():
                _validate_docx_archive(archive)
                return 'docx'
    except (BadZipFile, LargeZipFile):
        return None
    return None


def _validate_docx_archive(archive: ZipFile) -> None:
    infos = archive.infolist()
    if len(infos) > MAX_DOCX_ARCHIVE_MEMBERS:
        raise DocumentValidationError('DOCX 压缩包包含过多文件')

    total_uncompressed_bytes = 0
    for info in infos:
        if info.is_dir():
            continue
        if info.file_size > MAX_DOCX_MEMBER_BYTES:
            raise DocumentValidationError('DOCX 解压后单个文件过大')
        total_uncompressed_bytes += info.file_size
        if total_uncompressed_bytes > MAX_DOCX_UNCOMPRESSED_BYTES:
            raise DocumentValidationError('DOCX 解压后总大小超过限制')
        if info.file_size and (
            info.compress_size == 0
            or info.file_size / info.compress_size > MAX_DOCX_COMPRESSION_RATIO
        ):
            raise DocumentValidationError('DOCX 压缩比超过安全限制')


def _validate_source_path(path: Path) -> None:
    extension = path.suffix.lower()
    if extension not in {'.pdf', '.docx'}:
        raise DocumentValidationError('仅支持 PDF 和 DOCX 格式文件')
    size_bytes = path.stat().st_size
    if size_bytes > MAX_DOCUMENT_BYTES:
        raise DocumentValidationError('文件大小不能超过 100MB')
    with path.open('rb') as source_file:
        content = _read_limited_bytes(source_file)
    document_type = _detect_document_type(content)
    if document_type is None:
        raise DocumentValidationError('文件内容无法识别为 PDF 或 DOCX')
    if document_type != extension.lstrip('.'):
        raise DocumentValidationError('文件内容与扩展名不一致')


def _validate_document_page_limit(document_type: str, content: bytes) -> int:
    if document_type == 'pdf':
        try:
            pdf = fitz.open(stream=content, filetype='pdf')
        except (fitz.FileDataError, RuntimeError) as exc:
            raise DocumentValidationError('无法读取 PDF 文件') from exc
        with pdf:
            page_count = len(pdf)
        if page_count > MAX_DOCUMENT_PAGES:
            raise DocumentValidationError('PDF 文档不能超过 100 页')
        return page_count

    try:
        document = Document(BytesIO(content))
    except Exception as exc:  # python-docx raises several package-specific errors.
        raise DocumentValidationError('无法读取 DOCX 文件') from exc
    page_count = _docx_equivalent_page_count(document)
    if page_count > MAX_DOCUMENT_PAGES:
        raise DocumentValidationError('DOCX 文档等价页数不能超过 100 页')
    return page_count


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        pdf = fitz.open(path)
    except (fitz.FileDataError, RuntimeError) as exc:
        raise DocumentValidationError('无法读取 PDF 文件') from exc
    with pdf:
        page_count = len(pdf)
        if page_count > MAX_DOCUMENT_PAGES:
            raise DocumentValidationError('PDF 文档不能超过 100 页')
        page_details = [_pdf_page_detail(page, number + 1) for number, page in enumerate(pdf)]

    fragments = _pdf_fragments(page_details)
    if not fragments:
        raise DocumentValidationError('未找到可识别的编号题目')
    return ExtractedDocument(document_type='pdf', page_count=page_count, fragments=fragments)


def _pdf_page_detail(page, page_no: int) -> dict:
    blocks = sorted(page.get_text('blocks'), key=lambda block: (block[1], block[0]))
    text = '\n'.join(block[4] for block in blocks)
    starts = []
    for match in QUESTION_START.finditer(text):
        prefix = match.group(0).strip()
        rectangles = page.search_for(prefix)
        starts.append({
            'question_no': match.group(1),
            'offset': match.start(),
            'y': rectangles[0].y0 if rectangles else 0,
        })
    assets = []
    for image in page.get_images(full=True):
        xref = image[0]
        for rectangle in page.get_image_rects(xref):
            assets.append(DocumentAssetRef(
                reference=f'pdf:{xref}', page_no=page_no,
                bbox=(rectangle.x0, rectangle.y0, rectangle.x1, rectangle.y1),
            ))
    return {'page_no': page_no, 'text': text, 'blocks': blocks, 'starts': starts, 'assets': assets}


def _pdf_fragments(page_details: list[dict]) -> list[ExtractedQuestionFragment]:
    candidates = [
        {**start, 'page_index': index, 'page_no': detail['page_no']}
        for index, detail in enumerate(page_details)
        for start in detail['starts']
    ]
    fragments = []
    for index, candidate in enumerate(candidates):
        following = candidates[index + 1] if index + 1 < len(candidates) else None
        end_index = following['page_index'] if following else len(page_details) - 1
        text_parts = []
        for page_index in range(candidate['page_index'], end_index + 1):
            page_text = page_details[page_index]['text']
            start_offset = candidate['offset'] if page_index == candidate['page_index'] else 0
            end_offset = following['offset'] if following and page_index == end_index else len(page_text)
            text_parts.append(page_text[start_offset:end_offset].strip())
        fragments.append(ExtractedQuestionFragment(
            question_no=candidate['question_no'],
            text='\n'.join(part for part in text_parts if part),
            page_range=(candidate['page_no'], page_details[end_index]['page_no']),
            asset_refs=_pdf_assets_for_candidate(candidate, following, page_details),
        ))
    return fragments


def _pdf_assets_for_candidate(candidate: dict, following: dict | None, page_details: list[dict]):
    assets = []
    end_index = following['page_index'] if following else len(page_details) - 1
    for page_index in range(candidate['page_index'], end_index + 1):
        for asset in page_details[page_index]['assets']:
            y_center = (asset.bbox[1] + asset.bbox[3]) / 2 if asset.bbox else 0
            if page_index == candidate['page_index'] and y_center < candidate['y']:
                continue
            if following and page_index == following['page_index'] and y_center >= following['y']:
                continue
            assets.append(asset)
    return assets


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        document = Document(str(path))
    except Exception as exc:  # python-docx raises several package-specific errors.
        raise DocumentValidationError('无法读取 DOCX 文件') from exc

    image_refs = _docx_image_references(document)
    page_count = _docx_equivalent_page_count(document)
    if page_count > MAX_DOCUMENT_PAGES:
        raise DocumentValidationError('DOCX 文档等价页数不能超过 100 页')

    fragments = _docx_fragments(document, image_refs, max(page_count, 1))
    if not fragments:
        raise DocumentValidationError('未找到可识别的编号题目')
    return ExtractedDocument(document_type='docx', page_count=page_count, fragments=fragments)


def _docx_non_whitespace_characters(document) -> int:
    character_count, _table_count = _docx_body_content_metrics(document)
    return character_count


def _docx_equivalent_page_count(document) -> int:
    character_count, table_count = _docx_body_content_metrics(document)
    return (
        ceil(character_count / 1600)
        + ceil(table_count / 2)
        + _docx_image_occurrence_count(document)
    )


def _docx_body_content_metrics(document) -> tuple[int, int]:
    """Recursively count body text and tables without double-counting merged cells."""
    text_parts = [paragraph.text for paragraph in document.paragraphs]
    seen_cell_elements = set()
    seen_table_elements = set()
    table_count = 0

    def visit_table(table) -> None:
        nonlocal table_count
        table_element_id = id(table._tbl)
        if table_element_id in seen_table_elements:
            return
        seen_table_elements.add(table_element_id)
        table_count += 1
        for row in table.rows:
            for cell in row.cells:
                cell_element_id = id(cell._tc)
                if cell_element_id in seen_cell_elements:
                    continue
                seen_cell_elements.add(cell_element_id)
                text_parts.extend(paragraph.text for paragraph in cell.paragraphs)
                for nested_table in cell.tables:
                    visit_table(nested_table)

    for table in document.tables:
        visit_table(table)
    return len(re.sub(r'\s+', '', ''.join(text_parts))), table_count


def _docx_image_occurrence_count(document) -> int:
    """Count displayed image references, not unique relationship targets."""
    count = 0
    for part in document.part.package.parts:
        relations = getattr(part, 'rels', {})
        if not relations:
            continue
        for relation_id in DOCX_IMAGE_REFERENCE.findall(part.blob.decode('utf-8', errors='ignore')):
            relation = relations.get(relation_id)
            if relation and relation.reltype.endswith(RELATION_IMAGE_SUFFIX):
                count += 1
    return count


def _docx_image_references(document) -> dict[str, DocumentAssetRef]:
    return {
        relation_id: DocumentAssetRef(reference=f'docx:{relation.target_part.partname}')
        for relation_id, relation in document.part.rels.items()
        if relation.reltype.endswith(RELATION_IMAGE_SUFFIX)
    }


def _docx_fragments(document, image_refs: dict[str, DocumentAssetRef], page_count: int):
    fragments = []
    current = None
    pending_assets = []
    automatic_number_counters = {}
    for item in _docx_body_items(document):
        if isinstance(item, Paragraph):
            match = QUESTION_START.match(item.text)
            paragraph_assets = _paragraph_asset_refs(item, image_refs)
            automatic_question_no = _word_automatic_question_no(item, automatic_number_counters)
            if match or automatic_question_no:
                question_no = match.group(1) if match else automatic_question_no
                question_text = item.text[match.start():].strip() if match else (
                    f'{question_no}. {item.text.strip()}'
                )
                current = ExtractedQuestionFragment(
                    question_no=question_no, text=question_text,
                    page_range=(1, page_count), asset_refs=pending_assets + paragraph_assets,
                )
                fragments.append(current)
                pending_assets = []
            elif current is None:
                pending_assets.extend(paragraph_assets)
            else:
                if item.text.strip():
                    current.text = f'{current.text}\n{item.text.strip()}'.strip()
                current.asset_refs.extend(paragraph_assets)
        elif current is not None:
            current.tables.append([[cell.text for cell in row.cells] for row in item.rows])
    return fragments or _docx_unnumbered_fragments(document, image_refs, page_count)


def _word_automatic_question_no(paragraph, counters) -> str | None:
    """Read visible Word list numbering omitted from ``paragraph.text`` by python-docx."""
    properties = getattr(paragraph._p, 'pPr', None)
    number_properties = getattr(properties, 'numPr', None)
    if number_properties is None or not paragraph.text.strip():
        return None
    number_id = getattr(getattr(number_properties, 'numId', None), 'val', None)
    level = getattr(getattr(number_properties, 'ilvl', None), 'val', '0')
    if number_id in (None, '0', 0) or str(level) != '0':
        return None
    key = (str(number_id), str(level))
    counters[key] = counters.get(key, 0) + 1
    return str(counters[key])


def _docx_unnumbered_fragments(document, image_refs, page_count: int):
    """Split blank-separated practice groups only when all normal question markers are absent."""
    fragments = []
    text_parts = []
    asset_refs = []
    tables = []

    def finish_group():
        nonlocal text_parts, asset_refs, tables
        text = '\n'.join(part for part in text_parts if part).strip()
        if text and _looks_like_unnumbered_question(text):
            fragments.append(ExtractedQuestionFragment(
                question_no=str(len(fragments) + 1), text=text,
                tables=tables, page_range=(1, page_count), asset_refs=asset_refs,
            ))
        text_parts, asset_refs, tables = [], [], []

    for item in _docx_body_items(document):
        if isinstance(item, Paragraph):
            text = item.text.strip()
            paragraph_assets = _paragraph_asset_refs(item, image_refs)
            if not text and not paragraph_assets:
                finish_group()
                continue
            if text:
                text_parts.append(text)
            asset_refs.extend(paragraph_assets)
        else:
            tables.append([[cell.text for cell in row.cells] for row in item.rows])
    finish_group()
    return fragments


def _looks_like_unnumbered_question(text: str) -> bool:
    """Avoid sending plain titles and arbitrary prose to the per-question Qwen prompt."""
    return bool(OPTION_START.search(text) or QUESTION_SIGNAL.search(text))


def _docx_body_items(document):
    for child in document.element.body.iterchildren():
        if not isinstance(child.tag, str):
            continue
        if child.tag.endswith('}p'):
            yield Paragraph(child, document)
        elif child.tag.endswith('}tbl'):
            yield Table(child, document)


def _paragraph_asset_refs(paragraph, image_refs: dict[str, DocumentAssetRef]):
    relation_ids = DOCX_IMAGE_REFERENCE.findall(paragraph._element.xml)
    return [image_refs[relation_id] for relation_id in relation_ids if relation_id in image_refs]
